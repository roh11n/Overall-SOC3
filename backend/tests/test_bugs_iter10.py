"""Iteration 10 tests: real log-sources CSV (795 rows) validation + regressions."""
import os
import io
import requests
import pytest

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
ADMIN_EMAIL = "admin@mssp-soc.io"
ADMIN_PASSWORD = "Iris-df02a9d88045!Aa9"


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=30)
    assert r.status_code == 200, r.text
    return r.json().get("access_token") or r.json().get("token")


@pytest.fixture(scope="module")
def headers(token):
    return {"Authorization": f"Bearer {token}"}


# --- Log Sources summary ---
def test_logsources_summary(headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/log-sources",
                     params={"tenant_id": "all"}, headers=headers, timeout=30)
    assert r.status_code == 200, r.text
    data = r.json()
    print("LOGSOURCES:", data)
    # tolerate different envelope keys
    summary = data.get("summary", data)
    assert summary.get("total_log_sources") == 795
    assert summary.get("total_enabled_log_sources") == 795
    assert summary.get("log_sources_added") == 7
    assert summary.get("log_sources_added_total") == 794


# --- QRadar KPIs regression ---
def test_qradar_summary(headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/qradar",
                     params={"tenant_id": "all"}, headers=headers, timeout=30)
    assert r.status_code == 200, r.text
    print("QRADAR:", r.json())


# --- Executive dashboard regression ---
def test_executive(headers):
    r = requests.get(f"{BASE_URL}/api/dashboard/executive",
                     params={"tenant_id": "all"}, headers=headers, timeout=30)
    assert r.status_code == 200, r.text
    print("EXECUTIVE keys:", list(r.json().keys()))


# --- IRIS chat regression ---
def test_iris_chat(headers):
    r = requests.post(f"{BASE_URL}/api/copilot/chat",
                      json={"message": "hello", "tenant_id": "all"},
                      headers=headers, timeout=60)
    assert r.status_code == 200, r.text
    j = r.json()
    print("IRIS source:", j.get("source"))
    assert j.get("source") == "hf-llm"


# --- PPTX export ---
def test_pptx_export(headers, tmp_path):
    r = requests.get(f"{BASE_URL}/api/export/pptx",
                     params={"tenant_id": "all"}, headers=headers, timeout=120)
    assert r.status_code == 200, r.text[:500]
    assert r.headers.get("content-type", "").endswith("presentationml.presentation") or "officedocument" in r.headers.get("content-type", "")
    out = tmp_path / "out.pptx"
    out.write_bytes(r.content)
    from pptx import Presentation
    prs = Presentation(str(out))
    slides_text = []
    for s in prs.slides:
        txt = []
        for shape in s.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        txt.append(run.text)
        slides_text.append(" | ".join(t for t in txt if t))
    for i, t in enumerate(slides_text):
        print(f"SLIDE {i+1}: {t[:400]}")

    # Slide 3 (index 2): Executive Overview
    s3 = slides_text[2]
    assert "7" in s3, f"Expected new log sources '7' on slide 3, got: {s3}"
    assert "795" in s3, f"Expected 'log sources integrated' 795 on slide 3, got: {s3}"

    # Slide 4 (index 3): Executive Performance
    s4 = slides_text[3]
    assert "795" in s4
    assert "11.94" in s4, f"Expected MTTR 11.94 h on slide 4, got: {s4}"

    # Slide 8 (index 7): MITRE ATT&CK Hits by Tactic
    s8 = slides_text[7]
    assert "MITRE" in s8 or "ATT&CK" in s8 or "Tactic" in s8, f"MITRE chart title missing: {s8}"
